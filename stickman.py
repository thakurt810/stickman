"""
😂 Stickman Dad Joke Pipeline  — Full Scene Image Edition  v3
==============================================================
Portrait 720×1280  |  4 full-scene images  |  edge-tts audio
Jokes sourced from jokes.xlsx (columns: S.no, A, B, A again, Laugh)
State tracked in state.json (committed back to repo after each run)

Images used:
  black_talking.png  — black is speaking, orange listening
  orange_talking.png — orange is speaking, black listening
  left.png           — both laughing pose A
  right.png          — both laughing pose B  (alternated rapidly)

Run:
  pip install pillow numpy imageio[ffmpeg] edge-tts openpyxl
  python3 stickman_dadjoke_v2.py
"""

import os, json, asyncio, random, struct, wave, time, subprocess
import numpy as np
from datetime import datetime
from PIL import Image, ImageDraw, ImageFont

# ── CONFIGURATION ──────────────────────────────────────────────────────────
FPS            = 24
W, H           = 720, 1280
LAUGH_SWAP_FPS = 5
LAUGH_DURATION = 5.0

_SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
IMAGES_DIR   = os.environ.get("IMAGES_DIR",   os.path.join(_SCRIPT_DIR, "images"))
OUTPUT_DIR   = os.environ.get("OUTPUT_DIR",   os.path.join(_SCRIPT_DIR, "output"))
AUDIO_DIR    = os.environ.get("AUDIO_DIR",    os.path.join(_SCRIPT_DIR, "audio"))
METADATA_DIR = os.environ.get("METADATA_DIR", os.path.join(_SCRIPT_DIR, "metadata"))

# jokes.xlsx and state.json live in the repo root (same as script)
JOKES_FILE   = os.environ.get("JOKES_FILE",  os.path.join(_SCRIPT_DIR, "jokes.xlsx"))
STATE_FILE   = os.environ.get("STATE_FILE",  os.path.join(_SCRIPT_DIR, "state.json"))

GOOGLE_DRIVE_FOLDER   = os.environ.get("GOOGLE_DRIVE_FOLDER_ID", "")
METADATA_DRIVE_FOLDER = os.environ.get("METADATA_DRIVE_FOLDER_ID", "")
UPLOAD_TO_DRIVE       = os.environ.get("UPLOAD_TO_DRIVE", "false").lower() == "true"

VOICE_A = "en-US-GuyNeural"
VOICE_B = "en-US-AriaNeural"

MAX_RETRIES = 3

BUBBLE_BLACK_CX  = 210
BUBBLE_ORANGE_CX = 510
BUBBLE_TOP_Y     = 60
BUBBLE_AB_CX     = 360

# ── OPTIONAL IMPORTS ───────────────────────────────────────────────────────
try:
    import edge_tts
    HAS_EDGE_TTS = True
except ImportError:
    HAS_EDGE_TTS = False
    print("⚠️  edge-tts not installed — will use silent placeholders.")

try:
    import imageio.v3 as iio
    HAS_IMAGEIO = True
except ImportError:
    HAS_IMAGEIO = False
    print("⚠️  imageio not installed — pip install imageio[ffmpeg]")

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("⚠️  openpyxl not installed — pip install openpyxl")


# ══════════════════════════════════════════════════════════════════════════════
# 1.  JOKE LOADER  (from jokes.xlsx)
# ══════════════════════════════════════════════════════════════════════════════
def load_jokes() -> list:
    """
    Load all jokes from jokes.xlsx.
    Expected columns: S.no | A | B | A again | Laugh
    Returns list of dicts with keys: index, A, B, A_again, Laugh
    """
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required: pip install openpyxl")
    if not os.path.exists(JOKES_FILE):
        raise FileNotFoundError(f"jokes.xlsx not found at: {JOKES_FILE}")

    wb     = openpyxl.load_workbook(JOKES_FILE, data_only=True)
    ws     = wb.active
    rows   = list(ws.iter_rows(values_only=True))

    # Auto-detect header row
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    if "a" not in header:
        raise ValueError("Could not find column 'A' in jokes.xlsx header row.")

    # Map column names flexibly
    col = {
        "sno":     next((i for i, h in enumerate(header) if "s" in h and "no" in h), 0),
        "A":       next((i for i, h in enumerate(header) if h == "a"), 1),
        "B":       next((i for i, h in enumerate(header) if h == "b"), 2),
        "A_again": next((i for i, h in enumerate(header) if "again" in h), 3),
        "Laugh":   next((i for i, h in enumerate(header) if "laugh" in h), 4),
    }

    jokes = []
    for row_idx, row in enumerate(rows[1:], start=1):   # skip header
        a       = str(row[col["A"]]).strip()     if row[col["A"]]     else ""
        b       = str(row[col["B"]]).strip()     if row[col["B"]]     else ""
        a_again = str(row[col["A_again"]]).strip() if row[col["A_again"]] else ""
        laugh   = str(row[col["Laugh"]]).strip() if row[col["Laugh"]] else ""

        if a and b and a_again:   # skip empty rows
            jokes.append({
                "index":   row_idx,
                "A":       a,
                "B":       b,
                "A_again": a_again,
                "Laugh":   laugh if laugh else "Ha ha ha ha ha ha ha ha!",
            })

    print(f"  📋 Loaded {len(jokes)} jokes from jokes.xlsx")
    return jokes


# ══════════════════════════════════════════════════════════════════════════════
# 2.  STATE MANAGEMENT  (shuffle + pointer, persisted in state.json)
# ══════════════════════════════════════════════════════════════════════════════
def load_state(total_jokes: int) -> dict:
    """
    Load state.json.  If missing or stale (joke count changed), create fresh
    shuffled order.  Returns state dict with keys: order, pointer.
    """
    if os.path.exists(STATE_FILE):
        try:
            with open(STATE_FILE, "r") as f:
                state = json.load(f)
            # Validate: must have order covering all current jokes
            if (
                isinstance(state.get("order"), list)
                and len(state["order"]) == total_jokes
                and isinstance(state.get("pointer"), int)
            ):
                print(f"  📂 State loaded — pointer at {state['pointer']}/{total_jokes}")
                return state
            else:
                print("  ⚠️  State stale (joke count changed) — reshuffling.")
        except Exception as e:
            print(f"  ⚠️  Could not read state.json ({e}) — creating fresh state.")

    order = list(range(total_jokes))
    random.shuffle(order)
    state = {"order": order, "pointer": 0}
    print(f"  🔀 Fresh shuffle created ({total_jokes} jokes).")
    return state


def save_state(state: dict) -> None:
    with open(STATE_FILE, "w") as f:
        json.dump(state, f, indent=2)
    print(f"  💾 State saved (pointer={state['pointer']})")


def pick_joke(jokes: list, state: dict) -> dict:
    """
    Pick the next joke using the shuffled order.
    If we've used all jokes, reshuffle and start over.
    """
    pointer = state["pointer"]
    if pointer >= len(state["order"]):
        print("  🔄 All jokes used — reshuffling the deck!")
        order = list(range(len(jokes)))
        random.shuffle(order)
        state["order"]   = order
        state["pointer"] = 0
        pointer          = 0

    joke_index = state["order"][pointer]
    state["pointer"] = pointer + 1   # advance for next run

    chosen = jokes[joke_index]
    print(f"  🎯 Picked joke #{chosen['index']}: {chosen['A'][:50]}…")
    return chosen


def joke_to_dialogues(joke: dict) -> dict:
    """Convert joke dict to the dialogues list format used by the rest of pipeline."""
    return {
        "dialogues": [
            f"A: {joke['A']}",
            f"B: {joke['B']}",
            f"A: {joke['A_again']}",
            f"AB: {joke['Laugh']}",
        ]
    }


# ══════════════════════════════════════════════════════════════════════════════
# 3.  SCENE IMAGE LOADER
# ══════════════════════════════════════════════════════════════════════════════
def _load_scene(filename: str) -> Image.Image:
    path = os.path.join(IMAGES_DIR, filename)
    img  = Image.open(path).convert("RGB")
    img  = img.resize((W, H), Image.LANCZOS)
    return img

print("📂  Loading scene images…")
SCENES = {
    "black_talking":  _load_scene("black_talking.png"),
    "orange_talking": _load_scene("orange_talking.png"),
    "left":           _load_scene("left.png"),
    "right":          _load_scene("right.png"),
}
print("    ✅ All 4 scenes loaded.")


# ══════════════════════════════════════════════════════════════════════════════
# 4.  FONT HELPER
# ══════════════════════════════════════════════════════════════════════════════
def _get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    suffix = "-Bold" if bold else ""
    candidates = [
        f"/usr/share/fonts/truetype/dejavu/DejaVuSans{suffix}.ttf",
        f"/usr/share/fonts/truetype/liberation/LiberationSans{suffix}.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for p in candidates:
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size)
            except Exception:
                pass
    return ImageFont.load_default()


def _wrap(text: str, max_chars: int = 28) -> list:
    words = text.split()
    lines, line = [], []
    for w in words:
        if len(" ".join(line + [w])) <= max_chars:
            line.append(w)
        else:
            if line:
                lines.append(" ".join(line))
            line = [w]
    if line:
        lines.append(" ".join(line))
    return lines


# ══════════════════════════════════════════════════════════════════════════════
# 5.  SPEECH BUBBLE
# ══════════════════════════════════════════════════════════════════════════════
def draw_bubble(canvas: Image.Image, text: str, speaker: str, alpha: float) -> None:
    if alpha <= 0:
        return

    lines    = _wrap(text, max_chars=24)
    font     = _get_font(26)
    name_fnt = _get_font(22, bold=True)
    pad, lh  = 14, 32
    bh       = len(lines) * lh + pad * 2 + 36
    bw       = 300

    if speaker == "A":
        bcx        = BUBBLE_BLACK_CX
        tab_label  = "A"
        bg_col_rgb = (255, 255, 255)
    elif speaker == "B":
        bcx        = BUBBLE_ORANGE_CX
        tab_label  = "B"
        bg_col_rgb = (255, 255, 255)
    else:
        bcx        = BUBBLE_AB_CX
        tab_label  = "A & B 😂"
        bg_col_rgb = (255, 245, 180)

    bx    = bcx - bw // 2
    by    = BUBBLE_TOP_Y
    a_int = int(245 * min(1.0, alpha))
    layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d     = ImageDraw.Draw(layer)

    d.rounded_rectangle(
        [bx, by, bx + bw, by + bh],
        radius=18,
        fill=(*bg_col_rgb, a_int),
        outline=(60, 60, 60, a_int),
        width=2,
    )

    tab_w = max(70, len(tab_label) * 12 + 20)
    d.rounded_rectangle(
        [bx + 12, by - 28, bx + 12 + tab_w, by + 8],
        radius=9,
        fill=(40, 40, 40, a_int),
    )
    d.text(
        (bx + 12 + tab_w // 2, by - 10),
        tab_label,
        fill=(255, 255, 255, a_int),
        font=name_fnt,
        anchor="mm",
    )

    for i, ln in enumerate(lines):
        d.text(
            (bx + pad, by + pad + 10 + i * lh),
            ln,
            fill=(20, 20, 20, a_int),
            font=font,
        )

    canvas_rgba = canvas.convert("RGBA")
    canvas_rgba.alpha_composite(layer)
    canvas.paste(canvas_rgba.convert("RGB"))


# ══════════════════════════════════════════════════════════════════════════════
# 6.  FLOATING  HA!  TEXT
# ══════════════════════════════════════════════════════════════════════════════
HA_EMITTERS = [
    (110, 350, 0.9),
    (280, 280, 1.2),
    (390, 220, 0.8),
    (530, 300, 1.1),
    (650, 260, 1.0),
]

def draw_ha_text(canvas: Image.Image, t: float) -> None:
    font  = _get_font(44, bold=True)
    layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d     = ImageDraw.Draw(layer)

    for i, (ex, ey, spd) in enumerate(HA_EMITTERS):
        phase = (t * spd + i * 0.38) % 1.0
        y     = int(ey - 200 * phase)
        alpha = max(0, int(255 * (1.0 - phase)))
        if 0 < y < H and alpha > 15:
            colours = [(220,30,30),(180,0,0),(200,50,50),(230,60,0),(210,20,20)]
            col = (*colours[i % len(colours)], alpha)
            d.text((ex, y), "HA!", fill=col, font=font, anchor="mm")

    canvas_rgba = canvas.convert("RGBA")
    canvas_rgba.alpha_composite(layer)
    canvas.paste(canvas_rgba.convert("RGB"))


# ══════════════════════════════════════════════════════════════════════════════
# 7.  FRAME RENDERER
# ══════════════════════════════════════════════════════════════════════════════
def render_frame(scene_key, bubble_text, bubble_speaker, bubble_alpha, is_laugh, t):
    frame = SCENES[scene_key].copy()
    if is_laugh:
        draw_ha_text(frame, t)
    draw_bubble(frame, bubble_text, bubble_speaker, bubble_alpha)
    return np.array(frame)


# ══════════════════════════════════════════════════════════════════════════════
# 8.  AUDIO GENERATION
# ══════════════════════════════════════════════════════════════════════════════
async def _speak(text: str, voice: str, path: str) -> None:
    tts = edge_tts.Communicate(text, voice, rate="-20%")
    await tts.save(path)


def _generate_audio(text: str, voice: str, path: str) -> bool:
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            asyncio.run(_speak(text, voice, path))
            if os.path.exists(path) and os.path.getsize(path) > 0:
                return True
            raise RuntimeError("empty file")
        except Exception as e:
            print(f"      ⚠️  Audio attempt {attempt} failed: {e}")
            if attempt < MAX_RETRIES:
                time.sleep(1.5)
    return False


def _write_silent_wav(path: str, duration: float = 3.0, rate: int = 22050) -> None:
    n = int(rate * duration)
    with wave.open(path, "w") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(rate)
        wf.writeframes(struct.pack("<" + "h" * n, *([0] * n)))


def generate_all_audio(joke: dict) -> list:
    os.makedirs(AUDIO_DIR, exist_ok=True)
    audio_files = []

    for i, line in enumerate(joke["dialogues"]):
        speaker, text = line.split(":", 1)
        speaker = speaker.strip().upper()
        text    = text.strip()

        if speaker == "AB":
            tts_text = "Ha!  Ha!  Ha!  Ha!  Ha!  Ha!"
            voice    = VOICE_A
        elif speaker == "A":
            tts_text = text
            voice    = VOICE_A
        else:
            tts_text = text
            voice    = VOICE_B

        path = os.path.join(AUDIO_DIR, f"line_{i}.mp3")
        print(f"    🎙️  [{speaker}] {text[:55]}…")

        if HAS_EDGE_TTS:
            ok = _generate_audio(tts_text, voice, path)
            if not ok:
                print(f"    ⚠️  Audio failed for line {i+1} — using silence.")
                path = os.path.join(AUDIO_DIR, f"line_{i}.wav")
                _write_silent_wav(path, duration=4.0 if speaker == "AB" else 3.0)
        else:
            path = os.path.join(AUDIO_DIR, f"line_{i}.wav")
            _write_silent_wav(path, duration=4.0 if speaker == "AB" else 3.0)

        audio_files.append({"path": path, "speaker": speaker, "text": text})

    return audio_files


# ══════════════════════════════════════════════════════════════════════════════
# 9.  AUDIO DURATION
# ══════════════════════════════════════════════════════════════════════════════
def get_duration(path: str) -> float:
    try:
        r = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", path],
            capture_output=True, text=True, timeout=10,
        )
        return float(r.stdout.strip())
    except Exception:
        return 3.0


# ══════════════════════════════════════════════════════════════════════════════
# 10.  VIDEO ASSEMBLY
# ══════════════════════════════════════════════════════════════════════════════
def build_frames(joke: dict, audio_files: list) -> list:
    all_frames = []

    for i, line_data in enumerate(audio_files):
        speaker  = line_data["speaker"]
        text     = line_data["text"]
        path     = line_data["path"]
        is_laugh = speaker == "AB"

        duration = get_duration(path) if not is_laugh else LAUGH_DURATION
        n_frames = int(duration * FPS)

        print(f"  🎬 Scene {i+1}: [{speaker}] {text[:45]}… ({duration:.1f}s, {n_frames} frames)")

        for fi in range(n_frames):
            t = fi / FPS

            if is_laugh:
                swap_period = max(1, FPS // LAUGH_SWAP_FPS)
                scene_key   = "left" if (fi // swap_period) % 2 == 0 else "right"
            elif speaker == "A":
                scene_key = "black_talking"
            else:
                scene_key = "orange_talking"

            bubble_alpha = min(1.0, max(0.0, (t - 0.15) / 0.35))

            frame = render_frame(
                scene_key      = scene_key,
                bubble_text    = text,
                bubble_speaker = speaker,
                bubble_alpha   = bubble_alpha,
                is_laugh       = is_laugh,
                t              = t,
            )
            all_frames.append(frame)

    return all_frames


def write_silent_video(frames: list, out_path: str) -> None:
    print(f"  💾 Writing silent video ({len(frames)} frames)…")
    iio.imwrite(
        out_path,
        frames,
        fps=FPS,
        codec="libx264",
        quality=8,
        pixelformat="yuv420p",
    )


def merge_audio(video_path: str, audio_files: list, final_path: str) -> None:
    print("  🔊 Merging audio…")
    os.makedirs(AUDIO_DIR, exist_ok=True)

    list_path = os.path.join(AUDIO_DIR, "concat.txt")
    with open(list_path, "w") as f:
        for af in audio_files:
            abs_path = os.path.abspath(af["path"])
            f.write(f"file '{abs_path}'\n")

    concat_audio = os.path.join(AUDIO_DIR, "all_audio.aac")

    subprocess.run([
        "ffmpeg", "-y",
        "-f", "concat", "-safe", "0",
        "-i", list_path,
        "-c:a", "aac", "-b:a", "128k",
        concat_audio,
    ], capture_output=True, check=True)

    subprocess.run([
        "ffmpeg", "-y",
        "-i", video_path,
        "-i", concat_audio,
        "-c:v", "copy",
        "-c:a", "aac",
        "-shortest",
        final_path,
    ], capture_output=True, check=True)

    print(f"  ✅ Final video → {final_path}")


# ══════════════════════════════════════════════════════════════════════════════
# 11.  METADATA
# ══════════════════════════════════════════════════════════════════════════════
HASHTAGS = (
    "#Shorts #DadJokes #Stickman #FunnyAnimation #AnimatedShorts "
    "#DadJoke #ParkChat #StickFigure #FunnyVideos #KidsContent "
    "#FamilyFriendly #CartoonReels #JokeOfTheDay #CleanJokes "
    "#AnimationLovers #LaughOutLoud #FunnyReels #StickmanAnimation "
    "#DailyLaugh #FeelGoodVideos"
)

def generate_metadata(joke: dict) -> dict:
    setup     = joke["dialogues"][0].split(":", 1)[1].strip()
    punchline = joke["dialogues"][2].split(":", 1)[1].strip()
    desc = (
        f"😂 Two stickman friends share a dad joke in the park!\n\n"
        f"🎤 {setup}\n\n"
        f"🥁 {punchline}\n\n"
        f"{HASHTAGS}"
    )
    return {
        "youtube_title":       "😂 Stickman Dad Joke of the Day — Wait for the Punchline! #Shorts",
        "youtube_description": desc,
        "instagram_caption":   desc,
    }

def save_metadata(joke: dict, meta: dict, base_name: str) -> str:
    os.makedirs(METADATA_DIR, exist_ok=True)
    path = os.path.join(METADATA_DIR, f"{base_name}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump({
            "generated_at": datetime.utcnow().isoformat() + "Z",
            "joke_lines":   joke["dialogues"],
            **meta,
        }, f, ensure_ascii=False, indent=2)
    print(f"  📄 Metadata → {path}")
    return path


# ══════════════════════════════════════════════════════════════════════════════
# 12.  GOOGLE DRIVE UPLOAD  (optional)
# ══════════════════════════════════════════════════════════════════════════════
def _drive_upload(file_path: str, folder_id: str, mime: str = "video/mp4"):
    try:
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaFileUpload
        from google.auth.transport.requests import Request
        import pickle, base64
    except ImportError:
        print("  ⚠️  pip install google-api-python-client google-auth-oauthlib")
        return None

    creds = None
    token_b64 = os.environ.get("GDRIVE_TOKEN_BASE64", "")
    if token_b64:
        try:
            creds = pickle.loads(base64.b64decode(token_b64))
        except Exception as e:
            print(f"  ⚠️  Token decode failed: {e}")

    if creds is None and os.path.exists("token.pickle"):
        with open("token.pickle", "rb") as fh:
            creds = pickle.load(fh)

    if creds is None:
        print("  ❌ No Drive credentials. Skipping upload.")
        return None

    if not creds.valid and creds.expired and creds.refresh_token:
        try:
            creds.refresh(Request())
        except Exception as e:
            print(f"  ❌ Token refresh failed: {e}")
            return None

    try:
        svc   = build("drive", "v3", credentials=creds)
        meta  = {"name": os.path.basename(file_path),
                 "parents": [folder_id] if folder_id else []}
        media = MediaFileUpload(file_path, mimetype=mime, resumable=True)
        up    = svc.files().create(body=meta, media_body=media,
                                   fields="id,webViewLink").execute()
        link  = up.get("webViewLink")
        print(f"  ☁️  Uploaded → {link}")
        return link
    except Exception as e:
        print(f"  ❌ Drive upload failed: {e}")
        return None


# ══════════════════════════════════════════════════════════════════════════════
# 13.  MAIN
# ══════════════════════════════════════════════════════════════════════════════
def main():
    print("\n😂  Stickman Dad Joke Pipeline  v3  (Excel + State Tracking Edition)")
    print("=" * 70)

    if not HAS_IMAGEIO:
        print("❌  imageio required:  pip install imageio[ffmpeg]")
        return
    if not HAS_OPENPYXL:
        print("❌  openpyxl required:  pip install openpyxl")
        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    base_name = datetime.now().strftime("%Y%m%d_%H%M%S")
    print(f"   Base name: {base_name}")

    # ── Step 1: Load jokes from Excel ─────────────────────────────────────
    print("\n📖 Step 1 — Loading jokes from jokes.xlsx…")
    jokes = load_jokes()

    # ── Step 2: Pick next joke using state ────────────────────────────────
    print("\n🎯 Step 2 — Picking next joke…")
    state = load_state(len(jokes))
    chosen_joke = pick_joke(jokes, state)
    save_state(state)   # save immediately so even a crash won't repeat the joke

    joke = joke_to_dialogues(chosen_joke)
    print(f"   Joke lines:")
    for line in joke["dialogues"]:
        print(f"     {line}")

    # ── Step 3: Metadata ──────────────────────────────────────────────────
    print("\n🏷️  Step 3 — Metadata…")
    meta      = generate_metadata(joke)
    json_path = save_metadata(joke, meta, base_name)
    print(f"   Title: {meta['youtube_title']}")

    # ── Step 4: Audio ─────────────────────────────────────────────────────
    print("\n🎙️  Step 4 — Generating audio…")
    audio_files = generate_all_audio(joke)

    # ── Step 5: Render frames ─────────────────────────────────────────────
    print("\n🎬 Step 5 — Rendering frames…")
    frames = build_frames(joke, audio_files)
    print(f"   Total frames: {len(frames)}  ({len(frames)/FPS:.1f}s)")

    # ── Step 6: Write silent video ────────────────────────────────────────
    silent_path = os.path.join(OUTPUT_DIR, f"{base_name}_silent.mp4")
    write_silent_video(frames, silent_path)

    # ── Step 7: Mux audio ─────────────────────────────────────────────────
    final_path = os.path.join(OUTPUT_DIR, f"{base_name}.mp4")
    try:
        merge_audio(silent_path, audio_files, final_path)
        os.remove(silent_path)
    except Exception as e:
        print(f"  ⚠️  Audio mux failed ({e}) — keeping silent video.")
        final_path = silent_path

    # ── Step 8: Drive upload (optional) ───────────────────────────────────
    if UPLOAD_TO_DRIVE:
        print("\n☁️  Step 8 — Drive upload…")
        if GOOGLE_DRIVE_FOLDER:
            _drive_upload(final_path, GOOGLE_DRIVE_FOLDER, mime="video/mp4")
        if METADATA_DRIVE_FOLDER:
            _drive_upload(json_path, METADATA_DRIVE_FOLDER, mime="application/json")

    print(f"\n✅  Done!")
    print(f"   Video    → {final_path}")
    print(f"   Metadata → {json_path}\n")


if __name__ == "__main__":
    main()
